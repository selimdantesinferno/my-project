"""엑셀(.xlsx) / CSV 대량 업로드.

각 행마다: 글감 → AI 슬라이드 생성 → 렌더 → 예약 잡 등록.
scheduled_at 이 비어 있으면 즉시(지금) 발행 대상으로 등록됩니다.

지원 컬럼(헤더, 영문/한글 alias 허용):
  url, text, slide_count, template, tone, account,
  prompt, targets, first_comment, scheduled_at, provider
"""
from __future__ import annotations

import csv
import io

from . import ai, fetch, render, store, timeutil

# 헤더 정규화: 다양한 표기를 표준 키로 매핑
_ALIASES = {
    "url": "url", "링크": "url",
    "text": "text", "본문": "text", "글감": "text",
    "slide_count": "slide_count", "슬라이드수": "slide_count", "슬라이드": "slide_count",
    "template": "template", "템플릿": "template",
    "tone": "tone", "색상": "tone", "톤": "tone",
    "account": "account", "계정": "account", "계정명": "account",
    "prompt": "prompt", "프롬프트": "prompt",
    "targets": "targets", "플랫폼": "targets", "대상": "targets",
    "first_comment": "first_comment", "고정댓글": "first_comment", "첫댓글": "first_comment",
    "scheduled_at": "scheduled_at", "예약시각": "scheduled_at", "시간": "scheduled_at",
    "provider": "provider", "ai": "provider",
}

SAMPLE_HEADERS = [
    "url", "text", "slide_count", "template", "tone",
    "account", "prompt", "targets", "first_comment", "scheduled_at", "provider",
]


def _norm_key(key: str) -> str | None:
    k = (key or "").strip().lower().replace(" ", "")
    return _ALIASES.get(k) or _ALIASES.get(key.strip())


def _parse_rows(filename: str, content: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(c) if c is not None else "" for c in rows[0]]
        out = []
        for r in rows[1:]:
            if r is None or all(c is None for c in r):
                continue
            out.append({headers[i]: r[i] for i in range(min(len(headers), len(r)))})
        return out
    else:  # CSV
        text = content.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))


def _clean_row(raw: dict) -> dict:
    row: dict = {}
    for k, v in raw.items():
        std = _norm_key(str(k))
        if std and v is not None and str(v).strip() != "":
            row[std] = str(v).strip()
    return row


def _split_targets(value: str) -> list[str]:
    parts = [p.strip().lower() for p in value.replace(";", ",").split(",")]
    return [p for p in parts if p in ("instagram", "threads")]


async def process(filename: str, content: bytes) -> dict:
    """대량 파일을 처리하고 행별 결과 요약을 반환합니다."""
    raw_rows = _parse_rows(filename, content)
    results = []
    created = 0

    for idx, raw in enumerate(raw_rows, start=1):
        row = _clean_row(raw)
        try:
            source = row.get("text", "")
            if row.get("url"):
                fetched = await fetch.fetch_url_text(row["url"])
                source = f"{source}\n\n{fetched}".strip() if source else fetched
            if not source:
                raise ValueError("글감(text 또는 url)이 없습니다.")

            targets = _split_targets(row.get("targets", "instagram,threads"))
            if not targets:
                raise ValueError("targets 에 instagram/threads 가 없습니다.")

            data = await ai.generate_slides(
                source=source,
                slide_count=int(row.get("slide_count", 7)),
                account=row.get("account", ""),
                user_prompt=row.get("prompt", ""),
                provider=row.get("provider"),
            )
            images = render.render_deck(
                data["slides"], row.get("template", "toss"),
                row.get("tone", "blue"), row.get("account", ""),
            )
            scheduled_at = timeutil.to_utc_iso(row.get("scheduled_at"))
            job = await store.add_job(
                scheduled_at=scheduled_at, targets=targets, images=images,
                caption=data.get("caption", ""), hashtags=data.get("hashtags", []),
                first_comment=row.get("first_comment", ""),
                label=row.get("account", "") or f"행 {idx}",
            )
            created += 1
            results.append({"row": idx, "ok": True, "job_id": job["id"],
                            "scheduled_at": scheduled_at, "slides": len(images)})
        except Exception as e:  # noqa: BLE001 - 행 단위 격리
            results.append({"row": idx, "ok": False, "error": str(e)})

    return {"total": len(raw_rows), "created": created, "results": results}


def sample_csv() -> str:
    """대량 업로드 샘플 CSV 문자열."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(SAMPLE_HEADERS)
    w.writerow([
        "https://example.com/news", "", "7", "toss", "blue",
        "@my_account", "", "instagram,threads", "프로필 링크 확인!",
        "2026-06-01 09:00", "openai",
    ])
    w.writerow([
        "", "여기에 글감 본문을 직접 입력", "5", "magazine", "light",
        "@my_account", "통계 위주로 강조", "instagram", "",
        "", "gemini",
    ])
    return buf.getvalue()
